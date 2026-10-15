
import csv
import json
import logging
import sqlite3
from contextlib import closing
from datetime import datetime, timezone
from io import BytesIO, StringIO
from pathlib import Path
from uuid import uuid4

from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import ValidationError

from llm_service import generate_test_cases
from schemas import GenerateRequest, GenerateResponse, TestCase

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

app = FastAPI(
    title="AI-Powered Test Case Generator",
    version="1.2.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:3000",
        "http://127.0.0.1:3000",
        "http://localhost:5500",
        "http://127.0.0.1:5500",
        "http://localhost:5173",
        "http://127.0.0.1:5173",
    ],
    allow_credentials=False,
    allow_methods=["GET", "POST", "DELETE", "OPTIONS"],
    allow_headers=["Content-Type", "Authorization"],
)

# Store the database beside this Python file.
DATABASE_PATH = Path(__file__).resolve().parent / "test_cases.db"


def get_connection():
    connection = sqlite3.connect(str(DATABASE_PATH))
    connection.row_factory = sqlite3.Row
    return connection


def init_db():
    DATABASE_PATH.parent.mkdir(parents=True, exist_ok=True)

    with closing(get_connection()) as connection:
        with connection:
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS generations (
                    generation_id TEXT PRIMARY KEY,
                    requirement TEXT NOT NULL,
                    created_at TEXT NOT NULL,
                    test_cases_json TEXT NOT NULL
                )
                """
            )


# Create the table when the application starts importing this module.
init_db()


def save_generation(generation_id, requirement, test_cases):
    created_at = datetime.now(timezone.utc).isoformat()
    serialized_cases = json.dumps(
        [case.model_dump(mode="json") for case in test_cases]
    )

    with closing(get_connection()) as connection:
        with connection:
            connection.execute(
                """
                INSERT INTO generations
                    (generation_id, requirement, created_at, test_cases_json)
                VALUES (?, ?, ?, ?)
                """,
                (
                    generation_id,
                    requirement,
                    created_at,
                    serialized_cases,
                ),
            )


def fetch_generation(generation_id):
    with closing(get_connection()) as connection:
        row = connection.execute(
            """
            SELECT generation_id, requirement, created_at, test_cases_json
            FROM generations
            WHERE generation_id = ?
            """,
            (generation_id,),
        ).fetchone()

    return row


def get_saved_test_cases(row):
    raw_cases = json.loads(row["test_cases_json"])
    return [TestCase.model_validate(case) for case in raw_cases]


@app.get("/health")
def health_check():
    return {"status": "healthy"}


@app.post(
    "/api/generate-test-cases",
    response_model=GenerateResponse,
)
def create_test_cases(request: GenerateRequest):
    requirement = request.requirement.strip()

    if not requirement:
        raise HTTPException(
            status_code=422,
            detail="Requirement must not be blank.",
        )

    try:
        test_cases = generate_test_cases(requirement)

        if not 5 <= len(test_cases) <= 10:
            raise ValueError("Generation must produce 5 to 10 test cases.")

        generation_id = str(uuid4())

        # Save the requirement and generated cases before responding.
        save_generation(generation_id, requirement, test_cases)

        return GenerateResponse(
            generation_id=generation_id,
            test_cases=test_cases,
        )

    except ValidationError:
        logger.warning("Generated test-case data failed validation.")
        raise HTTPException(
            status_code=502,
            detail="The AI returned invalid test-case data. Please try again.",
        ) from None

    except Exception as exc:
        # Log the exception type only; never log credentials or API keys.
        logger.error(
            "Test-case generation or persistence failed (%s).",
            type(exc).__name__,
        )
        raise HTTPException(
            status_code=502,
            detail="Test-case generation failed. Please try again later.",
        ) from None


@app.get("/api/history")
def list_history():
    with closing(get_connection()) as connection:
        rows = connection.execute(
            """
            SELECT generation_id, requirement, created_at, test_cases_json
            FROM generations
            ORDER BY created_at DESC
            """
        ).fetchall()

    items = []

    for row in rows:
        cases = json.loads(row["test_cases_json"])
        items.append(
            {
                "id": row["generation_id"],
                "requirement": row["requirement"],
                "created_at": row["created_at"],
                "test_case_count": len(cases),
                "positive_count": sum(
                    case.get("test_type") == "Positive" for case in cases
                ),
                "negative_count": sum(
                    case.get("test_type") == "Negative" for case in cases
                ),
            }
        )

    return {"items": items}


@app.get("/api/history/{generation_id}")
def get_history_item(generation_id: str):
    row = fetch_generation(generation_id)

    if row is None:
        raise HTTPException(
            status_code=404,
            detail="History entry not found.",
        )

    return {
        "id": row["generation_id"],
        "requirement": row["requirement"],
        "created_at": row["created_at"],
        "test_cases": [case.model_dump(mode="json")
                       for case in get_saved_test_cases(row)],
    }


@app.delete("/api/history/{generation_id}", status_code=204)
def delete_history_item(generation_id: str):
    with closing(get_connection()) as connection:
        with connection:
            cursor = connection.execute(
                "DELETE FROM generations WHERE generation_id = ?",
                (generation_id,),
            )

    if cursor.rowcount == 0:
        raise HTTPException(
            status_code=404,
            detail="History entry not found.",
        )


@app.get("/api/export/{generation_id}")
def export_test_cases(generation_id: str, format: str = "xlsx"):
    if format not in {"csv", "json", "xlsx"}:
        raise HTTPException(
            status_code=400,
            detail="Unsupported export format. Use csv, json, or xlsx.",
        )

    row = fetch_generation(generation_id)

    if row is None:
        raise HTTPException(
            status_code=404,
            detail="Generated test cases were not found. Generate them again.",
        )

    test_cases = get_saved_test_cases(row)
    filename = f"test_cases_{generation_id}.{format}"

    if format == "json":
        content = json.dumps(
            [case.model_dump(mode="json") for case in test_cases],
            indent=2,
        ).encode("utf-8")
        media_type = "application/json"

    elif format == "csv":
        output = StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(
            [
                "Test Case ID",
                "Title",
                "Preconditions",
                "Test Steps",
                "Expected Result",
                "Priority",
                "Test Type",
            ]
        )

        for case in test_cases:
            writer.writerow(
                [
                    case.test_case_id,
                    case.title,
                    "\n".join(case.preconditions),
                    "\n".join(case.test_steps),
                    case.expected_result,
                    case.priority,
                    case.test_type,
                ]
            )

        content = output.getvalue().encode("utf-8-sig")
        media_type = "text/csv; charset=utf-8"

    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Test Cases"

        worksheet.append(
            [
                "Test Case ID",
                "Title",
                "Preconditions",
                "Test Steps",
                "Expected Result",
                "Priority",
                "Test Type",
            ]
        )

        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78",
            )

        for case in test_cases:
            worksheet.append(
                [
                    case.test_case_id,
                    case.title,
                    "\n".join(case.preconditions),
                    "\n".join(case.test_steps),
                    case.expected_result,
                    case.priority,
                    case.test_type,
                ]
            )

        for column_cells in worksheet.columns:
            column_letter = column_cells[0].column_letter
            max_length = max(
                len(str(cell.value or ""))
                for cell in column_cells
            )
            worksheet.column_dimensions[column_letter].width = min(
                max(max_length + 2, 15),
                45,
            )

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        output = BytesIO()
        workbook.save(output)
        content = output.getvalue()
        media_type = (
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )

    return StreamingResponse(
        BytesIO(content),
        media_type=media_type,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        },
    )
